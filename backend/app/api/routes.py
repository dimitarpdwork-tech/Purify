"""REST API."""
from __future__ import annotations

import csv
import datetime as dt
import io
import uuid
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from fastapi.security import OAuth2PasswordRequestForm
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import func, select, text, update
from sqlalchemy.orm import Session

from app.api import filters as flt
from app.api.serializers import EXPLORER_COLUMNS, website_row
from app.core.config import get_settings
from app.core.db import get_db
from app.core.security import create_access_token, get_current_user, hash_password, require_role, verify_password
from app.models import (
    AdInventory, AdsTxtFile, AdsTxtRecord, AuditLog, Contact, DetectedTechnology, ExclusionRule, Job, ManualOverride,
    OpportunityAssessment, Organization, Outreach, PopularityObservation, SavedView, Scan, Score, SellerEntity, StageStatus,
    TechnologyEvidence, User, Website, WebsiteClassification, WebsiteSource,
)

api = APIRouter(prefix="/api")
viewer = require_role("viewer")
analyst = require_role("analyst")
admin = require_role("admin")

OUTREACH_STAGES = ["Discovered", "Qualified", "Researching", "Ready to Contact", "Contacted", "In Discussion", "Testing",
                   "Signed", "Rejected", "Existing Partner"]
QUEUES = ["VERIFIED_HIGH_VALUE", "HIGH_POTENTIAL_VERIFY_TRAFFIC", "TECHNICAL_OPPORTUNITY", "NEEDS_REVIEW", "LOW_PRIORITY", "EXCLUDED"]


def audit(db: Session, user: User | None, action: str, entity: str, entity_id: Any = None, data: dict | None = None) -> None:
    db.add(AuditLog(user_id=user.id if user else None, action=action, entity=entity, entity_id=str(entity_id) if entity_id is not None else None, data=data))


# ------------------------------------------------------------------ auth & users
@api.post("/auth/login")
def login(form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == form.username.lower()))
    if not user or not user.is_active or not verify_password(form.password, user.password_hash):
        raise HTTPException(401, "Incorrect email or password")
    audit(db, user, "login", "user", user.id)
    db.commit()
    return {"access_token": create_access_token(user), "token_type": "bearer", "user": {"id": user.id, "email": user.email, "role": user.role}}


@api.get("/auth/me")
def me(user: User = Depends(viewer)):
    return {"id": user.id, "email": user.email, "role": user.role}


class UserIn(BaseModel):
    email: EmailStr
    password: str = Field(min_length=10)
    role: str = Field(pattern="^(viewer|analyst|admin)$")


@api.get("/users")
def list_users(db: Session = Depends(get_db), user: User = Depends(admin)):
    return [{"id": u.id, "email": u.email, "role": u.role, "is_active": u.is_active} for u in db.scalars(select(User).order_by(User.id))]


@api.post("/users")
def create_user(body: UserIn, db: Session = Depends(get_db), user: User = Depends(admin)):
    if db.scalar(select(User.id).where(User.email == body.email.lower())):
        raise HTTPException(409, "email already exists")
    u = User(email=body.email.lower(), password_hash=hash_password(body.password), role=body.role)
    db.add(u)
    audit(db, user, "create_user", "user", body.email, {"role": body.role})
    db.commit()
    return {"id": u.id}


class UserPatch(BaseModel):
    role: str | None = Field(default=None, pattern="^(viewer|analyst|admin)$")
    is_active: bool | None = None
    password: str | None = Field(default=None, min_length=10)


@api.patch("/users/{uid}")
def patch_user(uid: int, body: UserPatch, db: Session = Depends(get_db), user: User = Depends(admin)):
    u = db.get(User, uid) or (_ for _ in ()).throw(HTTPException(404, "not found"))
    if body.role:
        u.role = body.role
    if body.is_active is not None:
        u.is_active = body.is_active
    if body.password:
        u.password_hash = hash_password(body.password)
    audit(db, user, "update_user", "user", uid, body.model_dump(exclude={"password"}, exclude_none=True))
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ dashboard
@api.get("/dashboard")
def dashboard(db: Session = Depends(get_db), user: User = Depends(viewer)):
    q = lambda sql, **p: db.execute(text(sql), p).all()  # noqa: E731
    one = lambda sql, **p: db.execute(text(sql), p).scalar()  # noqa: E731
    stages = q("SELECT stage, status, count(*) FROM stage_status GROUP BY 1,2")
    stage_map: dict[str, dict[str, int]] = {}
    for st, status, n in stages:
        stage_map.setdefault(st, {})[status] = n
    week = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=7)
    return {
        "totals": {
            "websites": one("SELECT count(*) FROM websites"),
            "excluded": one("SELECT count(*) FROM websites WHERE excluded"),
            "classified": one("SELECT count(*) FROM websites WHERE primary_category IS NOT NULL"),
            "publishers": one("SELECT count(*) FROM websites WHERE is_publisher"),
            "adtech_analyzed": one("SELECT count(*) FROM websites WHERE last_analyzed_at IS NOT NULL"),
            "with_opportunity": one("SELECT count(*) FROM websites WHERE opportunity_types IS NOT NULL"),
            "high_potential": one("SELECT count(*) FROM websites WHERE queue IN ('HIGH_POTENTIAL_VERIFY_TRAFFIC','VERIFIED_HIGH_VALUE')"),
            "needs_review": one("SELECT count(*) FROM websites WHERE queue='NEEDS_REVIEW'"),
        },
        "queues": dict(q("SELECT coalesce(queue,'UNSCORED'), count(*) FROM websites GROUP BY 1")),
        "categories": dict(q("SELECT primary_category, count(*) FROM websites WHERE primary_category IS NOT NULL GROUP BY 1 ORDER BY 2 DESC")),
        "opportunities": dict(q("SELECT t, count(*) FROM websites, unnest(opportunity_types) t GROUP BY 1 ORDER BY 1")),
        "monetization": dict(q("SELECT monetization_type, count(*) FROM websites WHERE monetization_type IS NOT NULL GROUP BY 1 ORDER BY 2 DESC")),
        "crux_buckets": dict(q("SELECT crux_best_bucket::text, count(*) FROM websites WHERE crux_best_bucket IS NOT NULL GROUP BY crux_best_bucket ORDER BY crux_best_bucket")),
        "sources": dict(q("SELECT source_type, count(DISTINCT website_id) FROM website_sources GROUP BY 1")),
        "stages": stage_map,
        "jobs": [{"id": j.id, "kind": j.kind, "status": j.status, "progress": j.progress, "created_at": j.created_at.isoformat()}
                 for j in db.scalars(select(Job).order_by(Job.id.desc()).limit(8))],
        "new_promising": [website_row(w) for w in db.scalars(
            select(Website).where(Website.first_discovered_at >= week, Website.queue.in_(["HIGH_POTENTIAL_VERIFY_TRAFFIC", "TECHNICAL_OPPORTUNITY"]))
            .order_by(Website.score_total.desc().nulls_last()).limit(10))],
        "top_prospects": [website_row(w) for w in db.scalars(
            select(Website).where(Website.queue.in_(["VERIFIED_HIGH_VALUE", "HIGH_POTENTIAL_VERIFY_TRAFFIC"]))
            .order_by(Website.score_total.desc().nulls_last()).limit(15))],
        "recent_failures": [dict(r._mapping) for r in q(
            "SELECT s.stage, w.host, left(s.last_error, 200) AS error, s.updated_at FROM stage_status s JOIN websites w ON w.id=s.website_id "
            "WHERE s.status='FAILED' ORDER BY s.updated_at DESC LIMIT 10")],
        "traffic_note": "Traffic estimates are not collected in this build; popularity (CrUX/Tranco) ranks sites for manual traffic research.",
    }


# ------------------------------------------------------------------ explorer
@api.post("/websites/search")
def search_websites(f: flt.WebsiteFilter, page: int = 1, page_size: int = Query(50, le=500),
                    db: Session = Depends(get_db), user: User = Depends(viewer)):
    stmt = flt.apply(select(Website), f).offset((page - 1) * page_size).limit(page_size)
    rows = [website_row(w) for w in db.scalars(stmt)]
    return {"total": flt.count(db, f), "page": page, "page_size": page_size, "rows": rows, "columns": EXPLORER_COLUMNS}


class ManualAdd(BaseModel):
    domains: list[str] = Field(min_length=1, max_length=1000)
    note: str | None = None


@api.post("/websites")
def add_websites(body: ManualAdd, db: Session = Depends(get_db), user: User = Depends(analyst)):
    from app.services.normalize import try_normalize
    from app.services.registry import Discovered, register

    items, invalid = [], []
    for d in body.domains:
        n, err = try_normalize(d)
        if n:
            items.append(Discovered(norm=n, source_type="manual", source_ref=f"user:{user.id}", original=d,
                                    payload={"note": body.note} if body.note else None, priority=50))
        else:
            invalid.append({"value": d, "error": err})
    stats = register(db, items)
    audit(db, user, "add_websites", "website", None, {"count": len(items)})
    db.commit()
    ids = list(db.scalars(select(Website.id).where(Website.host.in_([i.norm.host for i in items]))))
    return {"added": stats.new_websites, "existing": stats.existing_websites, "invalid": invalid, "ids": ids}


def _w(db: Session, wid: int) -> Website:
    w = db.get(Website, wid)
    if not w:
        raise HTTPException(404, "website not found")
    return w


@api.get("/websites/by-host/{host}")
def website_by_host(host: str, db: Session = Depends(get_db), user: User = Depends(viewer)):
    from app.services.normalize import try_normalize

    n, _ = try_normalize(host)
    wid = db.scalar(select(Website.id).where(Website.host == (n.host if n else host)))
    if not wid:
        raise HTTPException(404, "website not found")
    return {"id": wid}


@api.get("/websites/{wid}")
def website_profile(wid: int, db: Session = Depends(get_db), user: User = Depends(viewer)):
    w = _w(db, wid)
    cls_all = db.scalars(select(WebsiteClassification).where(WebsiteClassification.website_id == wid).order_by(WebsiteClassification.id.desc()).limit(10)).all()
    dets = db.scalars(select(DetectedTechnology).where(DetectedTechnology.website_id == wid, DetectedTechnology.is_current.is_(True))
                      .order_by(DetectedTechnology.category, DetectedTechnology.vendor_key)).all()
    ev = {}
    if dets:
        for e in db.scalars(select(TechnologyEvidence).where(TechnologyEvidence.detection_id.in_([d.id for d in dets]))):
            ev.setdefault(e.detection_id, []).append({"kind": e.kind, "value": e.value, "page_url": e.page_url, "observation": e.observation, "scan_id": e.scan_id})
    ads = db.scalars(select(AdsTxtFile).where(AdsTxtFile.website_id == wid).order_by(AdsTxtFile.id.desc()).limit(5)).all()
    cur_ads = next((a for a in ads if a.is_current), None)
    records = []
    if cur_ads:
        recs = db.execute(select(AdsTxtRecord, SellerEntity).outerjoin(
            SellerEntity, (SellerEntity.ad_system_domain == AdsTxtRecord.ad_system_domain) & (SellerEntity.seller_id == AdsTxtRecord.seller_account_id))
            .where(AdsTxtRecord.file_id == cur_ads.id).order_by(AdsTxtRecord.line_no)).all()
        for r, se in recs:
            records.append({"ad_system_domain": r.ad_system_domain, "seller_account_id": r.seller_account_id, "relationship": r.relationship,
                            "cert_authority_id": r.cert_authority_id, "vendor_key": r.vendor_key, "line_no": r.line_no,
                            "seller": ({"name": se.name, "domain": se.domain, "seller_type": se.seller_type, "is_confidential": se.is_confidential,
                                        "domain_matches_site": bool(se.domain and se.domain.lower().endswith(w.registrable_domain))} if se else None)})
    opps = db.scalars(select(OpportunityAssessment).where(OpportunityAssessment.website_id == wid, OpportunityAssessment.is_current.is_(True))
                      .order_by(OpportunityAssessment.confidence.desc())).all()
    sc = db.scalar(select(Score).where(Score.website_id == wid, Score.is_current.is_(True)).order_by(Score.id.desc()).limit(1))
    score_hist = db.execute(select(Score.created_at, Score.total, Score.queue).where(Score.website_id == wid).order_by(Score.id.desc()).limit(20)).all()
    pops = db.scalars(select(PopularityObservation).where(PopularityObservation.website_id == wid).order_by(PopularityObservation.period)).all()
    inv = db.scalar(select(AdInventory).where(AdInventory.website_id == wid, AdInventory.is_current.is_(True)).order_by(AdInventory.id.desc()).limit(1))
    scans = db.scalars(select(Scan).where(Scan.website_id == wid).order_by(Scan.id.desc()).limit(15)).all()
    stages = db.scalars(select(StageStatus).where(StageStatus.website_id == wid)).all()
    sources = db.scalars(select(WebsiteSource).where(WebsiteSource.website_id == wid).order_by(WebsiteSource.first_seen_at)).all()
    overrides = db.scalars(select(ManualOverride).where(ManualOverride.website_id == wid).order_by(ManualOverride.id.desc())).all()
    outreach = db.get(Outreach, wid)
    org = db.get(Organization, w.organization_id) if w.organization_id else None
    contacts = db.scalars(select(Contact).where(Contact.website_id == wid)).all()
    stage_order = ["qualify", "classify", "supply", "adtech", "score"]
    return {
        "website": website_row(w) | {"preferred_origin": w.preferred_origin, "canonical_host": w.canonical_host,
                                     "redirects_to_website_id": w.redirects_to_website_id},
        "organization": {"name": org.name, "source": org.source} if org else None,
        "sources": [{"type": s.source_type, "ref": s.source_ref, "original": s.original_value, "payload": s.payload,
                     "first_seen": s.first_seen_at.isoformat(), "last_seen": s.last_seen_at.isoformat()} for s in sources],
        "stages": sorted([{"stage": s.stage, "status": s.status, "attempts": s.attempts, "last_error": s.last_error, "detail": s.detail,
                           "updated_at": s.updated_at.isoformat()} for s in stages], key=lambda x: stage_order.index(x["stage"]) if x["stage"] in stage_order else 99),
        "classification": [{"id": c.id, "primary_category": c.primary_category, "secondary": c.secondary_categories, "is_publisher": c.is_publisher,
                            "confidence": c.confidence, "supporting": c.supporting_evidence, "contradictory": c.contradictory_evidence,
                            "pages_inspected": c.pages_inspected, "features": c.features, "method": c.method, "engine": c.engine_version,
                            "is_current": c.is_current, "created_at": c.created_at.isoformat()} for c in cls_all],
        "popularity": [{"source": p.source, "period": p.period, "country": p.country, "rank_bucket": p.rank_bucket, "rank": p.rank,
                        "origin": p.origin, "source_url": p.source_url, "retrieved_at": p.retrieved_at.isoformat()} for p in pops],
        "technologies": [{"id": d.id, "vendor_key": d.vendor_key, "vendor_name": d.vendor_name, "category": d.category, "status": d.status,
                          "confidence": d.confidence, "detail": d.detail, "evidence": ev.get(d.id, []), "engine": d.engine_version,
                          "created_at": d.created_at.isoformat()} for d in dets],
        "inventory": inv.profile if inv else None,
        "ads_txt": {"files": [{"id": a.id, "url": a.url, "final_url": a.final_url, "http_status": a.http_status, "status": a.status,
                               "record_count": a.record_count, "malformed_count": a.malformed_count, "variables": a.variables,
                               "malformed_sample": a.malformed_sample, "fetched_at": a.fetched_at.isoformat(), "is_current": a.is_current} for a in ads],
                    "records": records},
        "opportunities": [{"type": o.opportunity_type, "label": o.label, "confidence": o.confidence, "evidence": o.evidence,
                           "hypothesis": o.hypothesis, "missing": o.missing_information, "verification": o.verification_step,
                           "approach": o.suggested_approach, "engine": o.engine_version, "created_at": o.created_at.isoformat()} for o in opps],
        "score": ({"popularity": sc.popularity, "ad_opportunity": sc.ad_opportunity, "audience_geo": sc.audience_geo, "inventory": sc.inventory,
                   "actionability": sc.actionability, "evidence_quality": sc.evidence_quality, "total": sc.total, "queue": sc.queue,
                   "computed_queue": sc.computed_queue, "evidence_confidence": sc.evidence_confidence, "explanation": sc.explanation,
                   "missing_critical": sc.missing_critical, "next_action": sc.next_action, "max": (sc.weights or {}).get("max"),
                   "engine": sc.engine_version, "created_at": sc.created_at.isoformat()} if sc else None),
        "score_history": [{"at": a.isoformat(), "total": t, "queue": q} for a, t, q in score_hist],
        "scans": [{"id": s.id, "kind": s.kind, "method": s.method, "status": s.status, "error": s.error, "started_at": s.started_at.isoformat(),
                   "pages": [{"id": p.id, "url": p.url, "final_url": p.final_url, "role": p.page_role, "http_status": p.http_status,
                              "blocked_reason": p.blocked_reason, "html_sha256": p.html_sha256, "network_request_count": len(p.network_requests or [])}
                             for p in s.pages]} for s in scans],
        "overrides": [{"field": o.field, "value": o.value, "reason": o.reason, "user_id": o.user_id, "active": o.active,
                       "created_at": o.created_at.isoformat()} for o in overrides],
        "outreach": {"stage": outreach.stage, "notes": outreach.notes, "owner_id": outreach.owner_id, "imported_status": outreach.imported_status} if outreach else None,
        "contacts": [{"email": c.email, "name": c.name, "role": c.role, "source_url": c.source_url, "verification_status": c.verification_status} for c in contacts],
        "business_links": next((((c.features or {}).get("home") or {}).get("business_links") for c in cls_all if c.is_current), None),
    }


@api.get("/scan-pages/{pid}/network")
def scan_page_network(pid: int, db: Session = Depends(get_db), user: User = Depends(viewer)):
    from app.models import ScanPage

    p = db.get(ScanPage, pid) or (_ for _ in ()).throw(HTTPException(404, "not found"))
    return {"url": p.url, "network_requests": p.network_requests or [], "script_urls": p.script_urls or [], "page_globals": p.page_globals}


class RescanIn(BaseModel):
    stages: list[str] = Field(default_factory=lambda: ["classify"])


@api.post("/websites/{wid}/rescan")
def rescan(wid: int, body: RescanIn, db: Session = Depends(get_db), user: User = Depends(analyst)):
    from app.pipeline.claims import STAGES, enqueue

    _w(db, wid)
    for st in body.stages:
        if st not in STAGES:
            raise HTTPException(400, f"unknown stage {st}")
        enqueue(db, wid, st, priority=1000)
    audit(db, user, "rescan", "website", wid, {"stages": body.stages})
    db.commit()
    return {"queued": body.stages}


class OverrideIn(BaseModel):
    field: str = Field(pattern="^(queue|excluded)$")
    value: str
    reason: str = Field(min_length=3)


@api.post("/websites/{wid}/override")
def override(wid: int, body: OverrideIn, db: Session = Depends(get_db), user: User = Depends(analyst)):
    w = _w(db, wid)
    if body.field == "queue" and body.value not in QUEUES + ["CLEAR"]:
        raise HTTPException(400, "invalid queue")
    db.execute(update(ManualOverride).where(ManualOverride.website_id == wid, ManualOverride.field == body.field).values(active=False))
    if body.value != "CLEAR":
        db.add(ManualOverride(website_id=wid, field=body.field, value=body.value, reason=body.reason, user_id=user.id))
    if body.field == "excluded":
        w.excluded = body.value.lower() == "true"
        w.exclusion_reason = f"manual: {body.reason}" if w.excluded else None
    audit(db, user, "override", "website", wid, body.model_dump())
    db.commit()
    from app.pipeline.stages import run_score_for

    run_score_for(wid)
    return {"ok": True}


class OutreachIn(BaseModel):
    stage: str | None = None
    notes: str | None = None
    owner_id: int | None = None


@api.patch("/websites/{wid}/outreach")
def set_outreach(wid: int, body: OutreachIn, db: Session = Depends(get_db), user: User = Depends(analyst)):
    w = _w(db, wid)
    if body.stage and body.stage not in OUTREACH_STAGES:
        raise HTTPException(400, "invalid stage")
    o = db.get(Outreach, wid) or Outreach(website_id=wid)
    if body.stage:
        o.stage = body.stage
        w.outreach_stage = body.stage
    if body.notes is not None:
        o.notes = body.notes
    if body.owner_id is not None:
        o.owner_id = body.owner_id
    db.merge(o)
    audit(db, user, "outreach", "website", wid, body.model_dump(exclude_none=True))
    db.commit()
    return {"ok": True}


class BulkIn(BaseModel):
    ids: list[int] = Field(min_length=1, max_length=5000)
    action: str = Field(pattern="^(rescan|exclude|include|outreach_stage)$")
    value: str | None = None
    reason: str | None = None


@api.post("/websites/bulk")
def bulk(body: BulkIn, db: Session = Depends(get_db), user: User = Depends(analyst)):
    from app.pipeline.claims import enqueue

    if body.action == "rescan":
        for i in body.ids:
            enqueue(db, i, body.value or "classify", priority=500)
    elif body.action in ("exclude", "include"):
        db.execute(update(Website).where(Website.id.in_(body.ids)).values(
            excluded=body.action == "exclude", exclusion_reason=(f"manual bulk: {body.reason or ''}" if body.action == "exclude" else None)))
        for i in body.ids:
            enqueue(db, i, "score", priority=500)
    elif body.action == "outreach_stage":
        if body.value not in OUTREACH_STAGES:
            raise HTTPException(400, "invalid stage")
        for i in body.ids:
            o = db.get(Outreach, i) or Outreach(website_id=i)
            o.stage = body.value
            db.merge(o)
        db.execute(update(Website).where(Website.id.in_(body.ids)).values(outreach_stage=body.value))
    audit(db, user, f"bulk_{body.action}", "website", None, {"count": len(body.ids), "value": body.value})
    db.commit()
    return {"ok": True, "count": len(body.ids)}


@api.get("/outreach/board")
def outreach_board(owner_id: int | None = None, db: Session = Depends(get_db), user: User = Depends(viewer)):
    stmt = select(Website).where(Website.outreach_stage != "Discovered")
    if owner_id:
        stmt = stmt.join(Outreach, Outreach.website_id == Website.id).where(Outreach.owner_id == owner_id)
    board: dict[str, list] = {s: [] for s in OUTREACH_STAGES}
    for w in db.scalars(stmt.order_by(Website.score_total.desc().nulls_last()).limit(2000)):
        board.setdefault(w.outreach_stage, []).append(website_row(w))
    # Qualified suggestions: high-potential queue still in Discovered
    board["Discovered"] = [website_row(w) for w in db.scalars(select(Website).where(
        Website.outreach_stage == "Discovered", Website.queue.in_(["VERIFIED_HIGH_VALUE", "HIGH_POTENTIAL_VERIFY_TRAFFIC", "TECHNICAL_OPPORTUNITY"]))
        .order_by(Website.score_total.desc().nulls_last()).limit(100))]
    return {"stages": OUTREACH_STAGES, "board": board}


# ------------------------------------------------------------------ export
class ExportIn(BaseModel):
    filter: flt.WebsiteFilter = Field(default_factory=flt.WebsiteFilter)
    format: str = Field(default="csv", pattern="^(csv|xlsx)$")
    columns: list[str] | None = None
    limit: int = Field(default=50000, le=200000)


@api.post("/export")
def export(body: ExportIn, db: Session = Depends(get_db), user: User = Depends(viewer)):
    cols = [c for c in (body.columns or EXPLORER_COLUMNS) if c in EXPLORER_COLUMNS]
    stmt = flt.apply(select(Website), body.filter).limit(body.limit)
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%d_%H%M")
    audit(db, user, "export", "website", None, {"format": body.format, "filter": body.filter.model_dump(exclude_none=True)})
    db.commit()

    def cell(v):
        if isinstance(v, list):
            return ", ".join(map(str, v))
        return "" if v is None else v

    note = "Popularity = CrUX bucket / Tranco rank (not visits). Ad-tech statuses: DETECTED_ACTIVE_SIGNALS is runtime evidence, not proof of auctions, fill or contracts."
    if body.format == "csv":
        def gen():
            buf = io.StringIO()
            wr = csv.writer(buf)
            wr.writerow(cols)
            yield buf.getvalue()
            for w in db.execute(stmt).scalars().yield_per(1000):
                buf.seek(0)
                buf.truncate()
                row = website_row(w)
                wr.writerow([cell(row[c]) for c in cols])
                yield buf.getvalue()
        return StreamingResponse(gen(), media_type="text/csv",
                                 headers={"Content-Disposition": f'attachment; filename="purify_publishers_{stamp}.csv"',
                                          "X-Data-Notes": note})
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("publishers")
    ws.append(cols)
    for w in db.execute(stmt).scalars().yield_per(1000):
        row = website_row(w)
        ws.append([cell(row[c]) for c in cols])
    ns = wb.create_sheet("notes")
    ns.append([note])
    ns.append([f"Exported {dt.datetime.now(dt.timezone.utc).isoformat()} by {user.email}"])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="purify_publishers_{stamp}.xlsx"'})


# ------------------------------------------------------------------ imports
def _upload_dir() -> Path:
    p = Path(get_settings().data_dir) / "uploads"
    p.mkdir(parents=True, exist_ok=True)
    return p


@api.post("/imports/upload")
async def upload(file: UploadFile = File(...), user: User = Depends(analyst)):
    from app.jobs.handlers import preview_file

    suffix = Path(file.filename or "upload.csv").suffix.lower()
    if suffix not in (".csv", ".txt", ".tsv", ".xlsx", ".xlsm"):
        raise HTTPException(400, "only CSV/TSV/TXT/XLSX files are supported")
    token = uuid.uuid4().hex
    dest = _upload_dir() / f"{token}{suffix}"
    size = 0
    with open(dest, "wb") as fh:  # stream to disk: large files never sit in memory
        while chunk := await file.read(1 << 20):
            size += len(chunk)
            if size > 1_000_000_000:
                dest.unlink(missing_ok=True)
                raise HTTPException(413, "file larger than 1 GB")
            fh.write(chunk)
    pv = preview_file(dest)
    return {"token": token, "filename": file.filename, "size": size, **pv}


class ImportStart(BaseModel):
    token: str = Field(pattern="^[0-9a-f]{32}$")
    mapping: dict[str, str]
    label: str | None = None
    honor_legacy_exclusions: bool = True


@api.post("/imports/start")
def start_import(body: ImportStart, db: Session = Depends(get_db), user: User = Depends(analyst)):
    from app.jobs.base import create_job
    from app.worker import run_job_task

    if "domain" not in body.mapping:
        raise HTTPException(400, "mapping must include a domain column")
    matches = list(_upload_dir().glob(f"{body.token}.*"))
    if not matches:
        raise HTTPException(404, "upload not found")
    jid = create_job("import_file", {"path": str(matches[0]), "mapping": body.mapping, "label": body.label,
                                     "honor_legacy_exclusions": body.honor_legacy_exclusions}, user.id)
    run_job_task.delay(jid)
    audit(db, user, "import", "job", jid, {"mapping": body.mapping})
    db.commit()
    return {"job_id": jid}


# ------------------------------------------------------------------ jobs / discovery
JOB_KINDS_USER = {"crux_ingest", "tranco_ingest", "sellers_discovery", "rescore_all", "requeue_stage"}


class JobIn(BaseModel):
    kind: str
    params: dict[str, Any] = Field(default_factory=dict)


@api.post("/jobs")
def create_job_ep(body: JobIn, db: Session = Depends(get_db), user: User = Depends(analyst)):
    from app.jobs.base import create_job
    from app.worker import run_job_task

    if body.kind not in JOB_KINDS_USER:
        raise HTTPException(400, f"kind must be one of {sorted(JOB_KINDS_USER)}")
    if body.kind == "requeue_stage" and body.params.get("stage") not in ("qualify", "classify", "supply", "adtech", "score"):
        raise HTTPException(400, "requeue_stage needs params.stage")
    jid = create_job(body.kind, body.params, user.id)
    run_job_task.delay(jid)
    audit(db, user, "create_job", "job", jid, body.model_dump())
    db.commit()
    return {"job_id": jid}


@api.get("/jobs")
def list_jobs(limit: int = 50, db: Session = Depends(get_db), user: User = Depends(viewer)):
    return [{"id": j.id, "kind": j.kind, "status": j.status, "params": j.params, "progress": j.progress, "error": (j.error or "")[:500],
             "created_at": j.created_at.isoformat(), "started_at": j.started_at.isoformat() if j.started_at else None,
             "finished_at": j.finished_at.isoformat() if j.finished_at else None,
             "heartbeat_at": j.heartbeat_at.isoformat() if j.heartbeat_at else None}
            for j in db.scalars(select(Job).order_by(Job.id.desc()).limit(limit))]


@api.post("/jobs/{jid}/cancel")
def cancel_job(jid: int, db: Session = Depends(get_db), user: User = Depends(analyst)):
    j = db.get(Job, jid) or (_ for _ in ()).throw(HTTPException(404, "not found"))
    j.cancel_requested = True
    if j.status == "QUEUED":
        j.status = "CANCELLED"
    audit(db, user, "cancel_job", "job", jid)
    db.commit()
    return {"ok": True}


@api.post("/jobs/{jid}/resume")
def resume_job(jid: int, db: Session = Depends(get_db), user: User = Depends(analyst)):
    from app.worker import run_job_task

    j = db.get(Job, jid) or (_ for _ in ()).throw(HTTPException(404, "not found"))
    if j.status not in ("FAILED", "CANCELLED"):
        raise HTTPException(400, "only failed or cancelled jobs can be resumed")
    j.status, j.cancel_requested, j.error = "QUEUED", False, None
    db.commit()
    run_job_task.delay(jid)
    return {"ok": True}


@api.get("/pipeline/status")
def pipeline_status(db: Session = Depends(get_db), user: User = Depends(viewer)):
    from app.services.settings_store import get_setting

    rows = db.execute(text("SELECT stage, status, count(*) FROM stage_status GROUP BY 1,2")).all()
    out: dict[str, dict[str, int]] = {}
    for st, status, n in rows:
        out.setdefault(st, {})[status] = n
    return {"stages": out, "dispatcher": get_setting("dispatcher")}


STAGE_LABELS = {"qualify": "Filter (popularity & exclusions)", "classify": "Classify website", "supply": "Read ads.txt",
                "adtech": "Browser ad-tech scan", "score": "Score opportunity"}
STALL_AFTER_MIN = 10


@api.get("/pipeline/progress")
def pipeline_progress(db: Session = Depends(get_db), user: User = Depends(viewer)):
    """Live processing view: is it working, per-step progress and speed, ETA, current and recent activity."""
    from app.services.settings_store import get_setting

    rows = db.execute(text("""
        SELECT stage,
               count(*) FILTER (WHERE status = 'PENDING')                       AS pending,
               count(*) FILTER (WHERE status = 'RUNNING')                       AS running,
               count(*) FILTER (WHERE status = 'DONE')                          AS done,
               count(*) FILTER (WHERE status = 'SKIPPED')                       AS skipped,
               count(*) FILTER (WHERE status = 'FAILED')                        AS failed,
               count(*) FILTER (WHERE status IN ('DONE','SKIPPED','FAILED')
                                AND updated_at > now() - interval '1 hour')     AS finished_1h,
               count(*) FILTER (WHERE status IN ('DONE','SKIPPED','FAILED')
                                AND updated_at > now() - interval '10 minutes') AS finished_10m,
               max(updated_at) FILTER (WHERE status IN ('DONE','SKIPPED','FAILED')) AS last_finished
        FROM stage_status GROUP BY stage""")).mappings().all()
    by = {r["stage"]: r for r in rows}
    order = ["qualify", "classify", "supply", "adtech", "score"]
    stages, last_activity, remaining_total = [], None, 0
    for st in order:
        r = by.get(st)
        if not r:
            stages.append({"stage": st, "label": STAGE_LABELS[st], "pending": 0, "running": 0, "done": 0, "skipped": 0, "failed": 0,
                           "rate_per_hour": 0, "eta_minutes": None})
            continue
        # speed: prefer the last 10 minutes (reacts quickly), fall back to the last hour
        rate = r["finished_10m"] * 6 if r["finished_10m"] >= 3 else r["finished_1h"]
        remaining = r["pending"] + r["running"]
        remaining_total += remaining
        eta = round(remaining / rate * 60) if (rate and remaining) else (0 if not remaining else None)
        if r["last_finished"] and (last_activity is None or r["last_finished"] > last_activity):
            last_activity = r["last_finished"]
        stages.append({"stage": st, "label": STAGE_LABELS[st], "pending": r["pending"], "running": r["running"], "done": r["done"],
                       "skipped": r["skipped"], "failed": r["failed"], "rate_per_hour": rate, "eta_minutes": eta})

    running_jobs = db.scalars(select(Job).where(Job.status.in_(("RUNNING", "QUEUED"))).order_by(Job.id)).all()
    for j in running_jobs:
        if j.heartbeat_at and (last_activity is None or j.heartbeat_at > last_activity):
            last_activity = j.heartbeat_at

    paused = bool(get_setting("dispatcher").get("paused"))
    now = dt.datetime.now(dt.timezone.utc)
    idle_min = (now - last_activity).total_seconds() / 60 if last_activity else None
    if paused:
        state, msg = "paused", "Processing is paused (Settings → dispatcher)."
    elif remaining_total == 0 and not running_jobs:
        state, msg = "idle", "Nothing waiting. Start a discovery job or add domains to begin."
    elif idle_min is None or idle_min > STALL_AFTER_MIN:
        state, msg = "stalled", (f"Work is waiting but nothing finished in the last {STALL_AFTER_MIN} minutes. "
                                 "Check that Docker is running and the laptop is awake (docker compose ps).")
    else:
        state, msg = "working", "Processing websites."

    # the slowest remaining step decides when the batch is done
    etas = [s["eta_minutes"] for s in stages if s["pending"] + s["running"] > 0]
    known = [e for e in etas if e is not None]
    overall_eta = max(known) if known else None
    # some steps have no speed yet (nothing finished there): the estimate is a lower bound
    eta_partial = bool(known) and len(known) < len(etas)

    active = db.execute(text("""
        SELECT w.id, w.host, s.stage, s.updated_at FROM stage_status s JOIN websites w ON w.id = s.website_id
        WHERE s.status = 'RUNNING' ORDER BY s.updated_at DESC LIMIT 12""")).mappings().all()
    recent = db.execute(text("""
        SELECT w.id, w.host, s.stage, s.status, s.updated_at, w.primary_category, w.queue, w.score_total
        FROM stage_status s JOIN websites w ON w.id = s.website_id
        WHERE s.status IN ('DONE','SKIPPED','FAILED') AND s.updated_at > now() - interval '1 day'
        ORDER BY s.updated_at DESC LIMIT 15""")).mappings().all()
    return {
        "now": now.isoformat(),
        "health": {"state": state, "message": msg, "last_activity_at": last_activity.isoformat() if last_activity else None},
        "remaining": remaining_total,
        "overall_eta_minutes": overall_eta,
        "overall_eta_partial": eta_partial,
        "finished_last_hour": sum(by[s]["finished_1h"] for s in by if s == "score"),
        "stages": stages,
        "jobs": [{"id": j.id, "kind": j.kind, "status": j.status, "progress": {k: v for k, v in (j.progress or {}).items() if k != "invalid_sample"},
                  "params": j.params, "started_at": j.started_at.isoformat() if j.started_at else None} for j in running_jobs],
        "active": [{"id": a["id"], "host": a["host"], "stage": a["stage"], "label": STAGE_LABELS.get(a["stage"], a["stage"]),
                    "since": a["updated_at"].isoformat()} for a in active],
        "recent": [{"id": r["id"], "host": r["host"], "stage": r["stage"], "label": STAGE_LABELS.get(r["stage"], r["stage"]), "status": r["status"],
                    "at": r["updated_at"].isoformat(), "category": r["primary_category"], "queue": r["queue"], "score": r["score_total"]}
                   for r in recent],
    }


# ------------------------------------------------------------------ supply chain
@api.get("/supply/search")
def supply_search(ad_system: str | None = None, seller_id: str | None = None, relationship: str | None = None, limit: int = 200,
                  db: Session = Depends(get_db), user: User = Depends(viewer)):
    if not ad_system and not seller_id:
        raise HTTPException(400, "ad_system or seller_id required")
    stmt = select(Website.id, Website.host, Website.primary_category, Website.score_total, AdsTxtRecord.ad_system_domain,
                  AdsTxtRecord.seller_account_id, AdsTxtRecord.relationship).join(Website, Website.id == AdsTxtRecord.website_id) \
        .where(AdsTxtRecord.is_current.is_(True))
    if ad_system:
        stmt = stmt.where(AdsTxtRecord.ad_system_domain == ad_system.lower())
    if seller_id:
        stmt = stmt.where(AdsTxtRecord.seller_account_id == seller_id)
    if relationship:
        stmt = stmt.where(AdsTxtRecord.relationship == relationship.upper())
    rows = db.execute(stmt.order_by(Website.score_total.desc().nulls_last()).limit(limit)).all()
    return [dict(r._mapping) for r in rows]


@api.get("/websites/{wid}/related")
def related_publishers(wid: int, db: Session = Depends(get_db), user: User = Depends(viewer)):
    """Websites sharing a DIRECT seller account. This is an ownership *hypothesis* only: shared
    accounts are common for networks, and reseller accounts are excluded."""
    w = _w(db, wid)
    rows = db.execute(text("""
        SELECT o.website_id, w2.host, r.ad_system_domain, r.seller_account_id, se.name AS seller_name, se.seller_type, se.domain AS seller_domain
        FROM ads_txt_records r
        JOIN ads_txt_records o ON o.ad_system_domain = r.ad_system_domain AND o.seller_account_id = r.seller_account_id
             AND o.relationship = 'DIRECT' AND o.is_current AND o.website_id <> r.website_id
        JOIN websites w2 ON w2.id = o.website_id
        LEFT JOIN seller_entities se ON se.ad_system_domain = r.ad_system_domain AND se.seller_id = r.seller_account_id
        WHERE r.website_id = :wid AND r.relationship = 'DIRECT' AND r.is_current
        LIMIT 500"""), {"wid": wid}).all()
    groups: dict[str, dict] = {}
    for r in rows:
        key = f"{r.ad_system_domain}|{r.seller_account_id}"
        g = groups.setdefault(key, {"ad_system_domain": r.ad_system_domain, "seller_account_id": r.seller_account_id,
                                    "seller_name": r.seller_name, "seller_type": r.seller_type, "seller_domain": r.seller_domain,
                                    "websites": []})
        g["websites"].append({"id": r.website_id, "host": r.host})
    out = []
    for g in groups.values():
        intermediary = g["seller_type"] in ("INTERMEDIARY",) or len(g["websites"]) > 50
        g["hypothesis"] = ("shared account of an intermediary/network - NOT evidence of common ownership" if intermediary
                           else "possible related publisher (shared DIRECT seller account) - verify")
        out.append(g)
    return {"host": w.host, "groups": sorted(out, key=lambda g: len(g["websites"]))}


@api.get("/supply/graph/{wid}")
def supply_graph(wid: int, db: Session = Depends(get_db), user: User = Depends(viewer)):
    w = _w(db, wid)
    recs = db.execute(select(AdsTxtRecord.ad_system_domain, AdsTxtRecord.relationship, AdsTxtRecord.vendor_key,
                             func.count(AdsTxtRecord.id)).where(AdsTxtRecord.website_id == wid, AdsTxtRecord.is_current.is_(True))
                      .group_by(AdsTxtRecord.ad_system_domain, AdsTxtRecord.relationship, AdsTxtRecord.vendor_key)).all()
    obs = {d.vendor_key: d.status for d in db.scalars(select(DetectedTechnology).where(
        DetectedTechnology.website_id == wid, DetectedTechnology.is_current.is_(True)))}
    nodes = [{"id": "site", "label": w.host, "type": "site"}]
    edges = []
    seen = set()
    for sysd, rel, vk, n in recs:
        if sysd not in seen:
            nodes.append({"id": sysd, "label": sysd, "type": "ad_system", "vendor_key": vk, "observed_status": obs.get(vk) if vk else None})
            seen.add(sysd)
        edges.append({"from": "site", "to": sysd, "relationship": rel, "accounts": n})
    return {"nodes": nodes, "edges": edges}


# ------------------------------------------------------------------ settings / exclusions / views
@api.get("/settings")
def get_settings_ep(user: User = Depends(viewer)):
    from app.services.settings_store import all_settings

    s = get_settings()
    return {"settings": all_settings(), "environment": {
        "playwright_enabled": s.playwright_enabled, "respect_robots_txt": s.respect_robots_txt, "user_agent": s.user_agent,
        "per_host_min_interval_seconds": s.per_host_min_interval_seconds, "playwright_concurrency": s.playwright_concurrency,
        "sellers_json_sources": s.sellers_json_source_list, "consent_mode": s.consent_mode,
        "not_in_this_build": ["traffic estimates (researched manually)", "HTTP Archive / BigQuery discovery", "contact discovery"]}}


class SettingIn(BaseModel):
    value: Any


@api.put("/settings/{key}")
def put_setting(key: str, body: SettingIn, db: Session = Depends(get_db), user: User = Depends(admin)):
    from app.services.settings_store import _defaults, set_setting

    if key not in _defaults():
        raise HTTPException(400, "unknown setting")
    val = set_setting(key, body.value, user.id)
    audit(db, user, "update_setting", "setting", key, {"value": body.value})
    db.commit()
    return {"key": key, "value": val}


class ExclusionIn(BaseModel):
    kind: str = Field(pattern="^(host|registrable|suffix|keyword)$")
    pattern: str = Field(min_length=2, max_length=253)
    reason: str = Field(min_length=3)


@api.get("/exclusions")
def list_exclusions(db: Session = Depends(get_db), user: User = Depends(viewer)):
    return [{"id": e.id, "kind": e.kind, "pattern": e.pattern, "reason": e.reason, "created_at": e.created_at.isoformat()}
            for e in db.scalars(select(ExclusionRule).order_by(ExclusionRule.id.desc()))]


@api.post("/exclusions")
def add_exclusion(body: ExclusionIn, db: Session = Depends(get_db), user: User = Depends(admin)):
    db.add(ExclusionRule(kind=body.kind, pattern=body.pattern.lower().strip(), reason=body.reason, created_by=user.id))
    audit(db, user, "add_exclusion", "exclusion", body.pattern, body.model_dump())
    db.commit()
    return {"ok": True}


@api.delete("/exclusions/{eid}")
def del_exclusion(eid: int, db: Session = Depends(get_db), user: User = Depends(admin)):
    e = db.get(ExclusionRule, eid) or (_ for _ in ()).throw(HTTPException(404, "not found"))
    audit(db, user, "delete_exclusion", "exclusion", eid, {"pattern": e.pattern})
    db.delete(e)
    db.commit()
    return {"ok": True}


class ViewIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    filters: dict[str, Any]
    columns: list[str] | None = None


@api.get("/views")
def list_views(db: Session = Depends(get_db), user: User = Depends(viewer)):
    return [{"id": v.id, "name": v.name, "filters": v.filters, "columns": v.columns}
            for v in db.scalars(select(SavedView).where(SavedView.user_id == user.id).order_by(SavedView.name))]


@api.post("/views")
def save_view(body: ViewIn, db: Session = Depends(get_db), user: User = Depends(viewer)):
    flt.WebsiteFilter(**body.filters)  # validate
    v = SavedView(user_id=user.id, name=body.name, filters=body.filters, columns=body.columns)
    db.add(v)
    db.commit()
    return {"id": v.id}


@api.delete("/views/{vid}")
def delete_view(vid: int, db: Session = Depends(get_db), user: User = Depends(viewer)):
    v = db.get(SavedView, vid)
    if not v or v.user_id != user.id:
        raise HTTPException(404, "not found")
    db.delete(v)
    db.commit()
    return {"ok": True}


@api.get("/vendors")
def vendors(user: User = Depends(viewer)):
    from app.adtech.registry import load_registry

    reg = load_registry()
    return {"version": reg.version, "vendors": [{"key": v.key, "name": v.name, "category": v.category, "ads_txt": v.ads_txt,
                                                  "bidders": v.bidders} for v in reg.vendors]}


@api.get("/audit")
def audit_log(limit: int = 200, db: Session = Depends(get_db), user: User = Depends(admin)):
    return [{"at": a.at.isoformat(), "user_id": a.user_id, "action": a.action, "entity": a.entity, "entity_id": a.entity_id, "data": a.data}
            for a in db.scalars(select(AuditLog).order_by(AuditLog.id.desc()).limit(limit))]
