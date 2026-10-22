"""Job handlers: discovery ingestion, file imports, sellers.json discovery, recomputation."""
from __future__ import annotations

import csv
import datetime as dt
import logging
import time
from pathlib import Path
from typing import Any, Iterator

from sqlalchemy import select, update
from sqlalchemy.dialects.postgresql import insert

from app.core.config import get_settings
from app.core.db import session_scope
from app.discovery import datasets
from app.jobs.base import JobContext, handler
from app.models import Outreach, SellerEntity, SellersJsonFile, StageStatus, Website
from app.services.normalize import try_normalize
from app.services.registry import Discovered, register

log = logging.getLogger(__name__)
BATCH = 2000

LEGACY_LABELS = {
    # historical label -> (priority boost, note). Never treated as traffic or verified quality.
    "PRIORITY": (30, "legacy PRIORITY label"),
    "POSSIBLE": (15, "legacy POSSIBLE label"),
    "COMMERCIAL_REVIEW": (10, "legacy COMMERCIAL_REVIEW label"),
    "UNKNOWN": (12, "legacy UNKNOWN label - rescue candidate"),
    "BRANDED": (5, "legacy BRANDED label"),
    "EXCLUDED": (0, "legacy EXCLUDED label"),
}


def _register_with_retry(batch: list[Discovered]):
    """Registration is idempotent, so a batch that hits a deadlock / serialization failure is retried."""
    from sqlalchemy.exc import OperationalError

    for attempt in range(4):
        try:
            with session_scope() as s:
                return register(s, batch)
        except OperationalError as exc:
            if attempt == 3 or not any(k in str(exc).lower() for k in ("deadlock", "could not serialize")):
                raise
            time.sleep(0.5 * (attempt + 1))


def _flush(ctx: JobContext, batch: list[Discovered], row_no: int, **extra: Any) -> None:
    if batch:
        st = _register_with_retry(batch)
        ctx.incr(registered=st.received, new_websites=st.new_websites, existing_websites=st.existing_websites,
                 duplicates=st.duplicates_in_batch)
    ctx.save(checkpoint={"row": row_no}, **extra)


# ------------------------------------------------------------------ CrUX
@handler("crux_ingest")
def crux_ingest(ctx: JobContext) -> None:
    """params: country?, max_bucket, exclude_top? (skip sites in this band or higher, e.g. 10000 = start after
    the top 10K), limit? (stop after N websites, most popular bands first), requested? (started by a user)."""
    country = ctx.params.get("country") or None
    max_bucket = int(ctx.params.get("max_bucket") or get_settings().crux_max_rank_bucket)
    exclude_top = int(ctx.params.get("exclude_top") or 0)
    limit = int(ctx.params.get("limit") or 0)
    requested = bool(ctx.params.get("requested"))
    period = ctx.params.get("period") or ctx.checkpoint.get("period") or datasets.crux_latest_period(country)
    path, url = datasets.fetch_crux(period, country)
    ctx.save(checkpoint={"period": period}, source_url=url)
    start = int(ctx.checkpoint.get("row", 0))
    accepted = int(ctx.progress.get("accepted", 0))
    batch: list[Discovered] = []
    last = start
    ref = f"crux:{country or 'global'}:{period}"
    for row_no, origin, bucket in datasets.iter_crux(path, max_bucket):
        if row_no <= start:
            continue
        last = row_no
        if bucket <= exclude_top:
            continue
        if limit and accepted >= limit:
            break
        norm, err = try_normalize(origin)
        if not norm:
            ctx.incr(invalid=1)
            continue
        accepted += 1
        payload = {"bucket": bucket, "period": period, "country": country or "global"}
        if requested:
            payload["requested"] = True  # user chose this range: skip the automatic popularity gate
        batch.append(Discovered(
            norm=norm, source_type="crux", source_ref=ref, original=origin, payload=payload,
            priority=datasets.popularity_priority(crux_bucket=bucket),
            popularity=[{"source": "crux", "period": period, "country": (country or "global").upper() if country else "global",
                         "rank_bucket": bucket, "origin": origin, "source_url": url}],
        ))
        if len(batch) >= BATCH:
            _flush(ctx, batch, last, accepted=accepted)
            batch = []
    _flush(ctx, batch, last, finished_rows=last, accepted=accepted)


# ------------------------------------------------------------------ Tranco
@handler("tranco_ingest")
def tranco_ingest(ctx: JobContext) -> None:
    max_rank = int(ctx.params.get("max_rank") or get_settings().tranco_max_rank)
    if ctx.checkpoint.get("path"):
        path, url, list_id = Path(ctx.checkpoint["path"]), ctx.checkpoint["url"], ctx.checkpoint["list_id"]
    else:
        path, url, list_id = datasets.fetch_tranco()
        ctx.save(checkpoint={"path": str(path), "url": url, "list_id": list_id})
    start = int(ctx.checkpoint.get("row", 0))
    batch: list[Discovered] = []
    last = start
    for row_no, domain, rank in datasets.iter_tranco(path, max_rank):
        if row_no <= start:
            continue
        last = row_no
        norm, _ = try_normalize(domain)
        if not norm:
            ctx.incr(invalid=1)
            continue
        batch.append(Discovered(
            norm=norm, source_type="tranco", source_ref=f"tranco:{list_id}", original=domain,
            payload={"rank": rank, "list_id": list_id, **({"requested": True} if ctx.params.get("requested") else {})},
            priority=datasets.popularity_priority(tranco_rank=rank),
            popularity=[{"source": "tranco", "period": list_id, "rank": rank, "source_url": url}],
        ))
        if len(batch) >= BATCH:
            _flush(ctx, batch, last)
            batch = []
    _flush(ctx, batch, last, finished_rows=last)


# ------------------------------------------------------------------ CSV / XLSX import
def _iter_rows(path: Path) -> Iterator[tuple[int, dict[str, str]]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(next(rows, []) or [])]
        for i, row in enumerate(rows, start=1):
            yield i, {header[j]: ("" if v is None else str(v)) for j, v in enumerate(row) if j < len(header)}
        wb.close()
        return
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        sample = fh.read(65536)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(fh, dialect=dialect)
        for i, row in enumerate(reader, start=1):
            yield i, {k.strip() if k else "": (v or "") for k, v in row.items()}


def preview_file(path: Path, limit: int = 20) -> dict[str, Any]:
    rows = []
    for i, r in _iter_rows(path):
        rows.append(r)
        if i >= limit:
            break
    columns = list(rows[0].keys()) if rows else []
    guess = {}
    for c in columns:
        lc = c.lower()
        if "domain" in lc or "website" in lc or "url" in lc or lc in ("site", "host"):
            guess.setdefault("domain", c)
        elif "categ" in lc or "class" in lc or "label" in lc:
            guess.setdefault("legacy_label", c)
        elif "note" in lc or "comment" in lc:
            guess.setdefault("notes", c)
        elif "geo" in lc or "country" in lc:
            guess.setdefault("geo", c)
        elif "outreach" in lc or "status" in lc or "stage" in lc:
            guess.setdefault("outreach_status", c)
        elif "source" in lc:
            guess.setdefault("source", c)
    return {"columns": columns, "rows": rows, "suggested_mapping": guess}


@handler("import_file")
def import_file(ctx: JobContext) -> None:
    """params: path, mapping {domain, legacy_label?, category?, notes?, geo?, source?, outreach_status?}, label."""
    path = Path(ctx.params["path"])
    mapping: dict[str, str] = ctx.params["mapping"]
    label = ctx.params.get("label") or path.name
    dcol = mapping["domain"]
    start = int(ctx.checkpoint.get("row", 0))
    batch: list[Discovered] = []
    outreach: list[tuple[str, str]] = []
    excluded_hosts: list[tuple[str, str]] = []
    invalid_sample: list[dict] = list(ctx.progress.get("invalid_sample", []))
    last = start
    for row_no, row in _iter_rows(path):
        if row_no <= start:
            continue
        last = row_no
        raw = (row.get(dcol) or "").strip()
        norm, err = try_normalize(raw)
        if not norm:
            ctx.incr(invalid=1)
            if len(invalid_sample) < 50:
                invalid_sample.append({"row": row_no, "value": raw[:200], "error": err})
            continue
        payload = {k: row.get(v) for k, v in mapping.items() if k != "domain" and v and row.get(v) not in (None, "")}
        legacy = (payload.get("legacy_label") or "").strip().upper().replace(" ", "_")
        boost = LEGACY_LABELS.get(legacy, (8, None))[0]
        if legacy:
            payload["legacy_label"] = legacy
        batch.append(Discovered(norm=norm, source_type="legacy" if legacy else "csv", source_ref=f"import:{ctx.job_id}:{label}",
                                original=raw, payload=payload or None, priority=20 + boost))
        if payload.get("outreach_status"):
            outreach.append((norm.host, str(payload["outreach_status"])[:100]))
        if legacy == "EXCLUDED" and ctx.params.get("honor_legacy_exclusions", True):
            excluded_hosts.append((norm.host, "legacy label EXCLUDED (imported)"))
        if len(batch) >= BATCH:
            _flush(ctx, batch, last, invalid_sample=invalid_sample)
            _apply_side_tables(outreach, excluded_hosts)
            batch, outreach, excluded_hosts = [], [], []
    _flush(ctx, batch, last, invalid_sample=invalid_sample, finished_rows=last)
    _apply_side_tables(outreach, excluded_hosts)


def _apply_side_tables(outreach: list[tuple[str, str]], excluded: list[tuple[str, str]]) -> None:
    if not outreach and not excluded:
        return
    with session_scope() as s:
        hosts = list({h for h, _ in outreach} | {h for h, _ in excluded})
        ids = dict(s.execute(select(Website.host, Website.id).where(Website.host.in_(hosts))).all())
        if outreach:
            latest = {h: v for h, v in outreach if h in ids}  # same host may appear twice in a batch (www / case variants)
            stmt = insert(Outreach).values([{"website_id": ids[h], "imported_status": v, "stage": "Discovered"} for h, v in latest.items()])
            s.execute(stmt.on_conflict_do_update(index_elements=["website_id"], set_={"imported_status": stmt.excluded.imported_status}))
        for h, reason in excluded:
            if h in ids:
                s.execute(update(Website).where(Website.id == ids[h]).values(excluded=True, exclusion_reason=reason))


# ------------------------------------------------------------------ sellers.json discovery
@handler("sellers_discovery")
def sellers_discovery(ctx: JobContext) -> None:
    """Download sellers.json files of configured ad systems; store entities and register
    non-confidential PUBLISHER/BOTH domains as discovery leads."""
    from app.supply.sellersjson import iter_sellers_from_file

    sources = ctx.params.get("sources") or get_settings().sellers_json_source_list
    done = set(ctx.checkpoint.get("done", []))
    cache = Path(get_settings().data_dir) / "datasets"
    cache.mkdir(parents=True, exist_ok=True)
    for entry in sources:
        if entry in done:
            continue
        explicit, url = get_settings().split_sellers_source(entry)
        norm, _ = try_normalize(url)
        ad_system = explicit or (norm.registrable_domain if norm else url)
        dest = cache / f"sellers_{ad_system}_{dt.date.today().isoformat()}.json"
        status, count, pubs = "OK", 0, 0
        try:
            if not dest.exists():
                datasets.download(url, dest, max_bytes=get_settings().sellers_json_max_bytes)
            ents: list[dict] = []
            leads: list[Discovered] = []
            for e in iter_sellers_from_file(str(dest)):
                count += 1
                ents.append({**e, "ad_system_domain": ad_system})
                if e["domain"] and not e["is_confidential"] and e["seller_type"] in ("PUBLISHER", "BOTH"):
                    n, _ = try_normalize(e["domain"])
                    if n:
                        pubs += 1
                        leads.append(Discovered(norm=n, source_type="sellers_json", source_ref=f"sellers:{ad_system}",
                                                original=e["domain"], payload={"ad_system": ad_system, "seller_id": e["seller_id"],
                                                                               "name": e["name"], "seller_type": e["seller_type"]},
                                                priority=10))
                if len(ents) >= BATCH:
                    _store_sellers(ents)
                    ents = []
                if len(leads) >= BATCH:
                    st = _register_with_retry(leads)
                    ctx.incr(registered=st.received, new_websites=st.new_websites)
                    leads = []
            _store_sellers(ents)
            if leads:
                st = _register_with_retry(leads)
                ctx.incr(registered=st.received, new_websites=st.new_websites)
        except Exception as exc:  # one failing SSP must not abort the others
            status = f"FAILED: {type(exc).__name__}: {str(exc)[:200]}"
        with session_scope() as s:
            s.add(SellersJsonFile(ad_system_domain=ad_system, url=url, status=status[:20] if status == "OK" else "FAILED",
                                  seller_count=count, publisher_domain_count=pubs, meta={"detail": status}))
        done.add(entry)
        ctx.save(checkpoint={"done": sorted(done)}, **{f"sellers:{ad_system}": {"status": status, "sellers": count, "publisher_domains": pubs}})


def _store_sellers(ents: list[dict]) -> None:
    if not ents:
        return
    dedup = {(e["ad_system_domain"], e["seller_id"]): e for e in ents}
    now = dt.datetime.now(dt.timezone.utc)
    with session_scope() as s:
        stmt = insert(SellerEntity).values([{**e, "fetched_at": now} for e in dedup.values()])
        s.execute(stmt.on_conflict_do_update(
            index_elements=["ad_system_domain", "seller_id"],
            set_={"name": stmt.excluded.name, "domain": stmt.excluded.domain, "seller_type": stmt.excluded.seller_type,
                  "is_confidential": stmt.excluded.is_confidential, "fetched_at": now}))


# ------------------------------------------------------------------ recomputation without re-crawling
@handler("rescore_all")
def rescore_all(ctx: JobContext) -> None:
    """Re-run opportunity + scoring from stored observations for every analysed website."""
    from app.pipeline.stages import run_score_for

    last_id = int(ctx.checkpoint.get("last_id", 0))
    while True:
        with session_scope() as s:
            ids = list(s.scalars(select(StageStatus.website_id).where(
                StageStatus.stage == "score", StageStatus.status == "DONE", StageStatus.website_id > last_id)
                .order_by(StageStatus.website_id).limit(500)))
        if not ids:
            break
        for wid in ids:
            run_score_for(wid)
        last_id = ids[-1]
        ctx.incr(rescored=len(ids))
        ctx.save(checkpoint={"last_id": last_id})


@handler("requeue_stage")
def requeue_stage(ctx: JobContext) -> None:
    """Re-run a stage for websites matching a status (e.g. reclassify UNKNOWN / failed)."""
    stage = ctx.params["stage"]
    statuses = ctx.params.get("statuses", ["FAILED"])
    with session_scope() as s:
        n = s.execute(update(StageStatus).where(StageStatus.stage == stage, StageStatus.status.in_(statuses))
                      .values(status="PENDING", attempts=0, lease_until=None)).rowcount
    ctx.save(requeued=n)
