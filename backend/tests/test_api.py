"""API tests (auth, roles, explorer, profile, export integrity, imports, jobs) against the test DB."""
import csv
import io

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(clean_db, monkeypatch):
    monkeypatch.setenv("BOOTSTRAP_ADMIN_EMAIL", "admin@test.io")
    monkeypatch.setenv("BOOTSTRAP_ADMIN_PASSWORD", "admin-password-1")
    from app.core.config import get_settings

    get_settings.cache_clear()
    import app.main as main

    main.settings = get_settings()
    # avoid sending Celery tasks to a broker during API tests
    import app.worker as worker

    monkeypatch.setattr(worker.run_job_task, "delay", lambda *a, **k: None)
    with TestClient(main.app) as c:
        yield c
    get_settings.cache_clear()


def login(c, email="admin@test.io", pw="admin-password-1"):
    r = c.post("/api/auth/login", data={"username": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_auth_required_and_bad_token(client):
    assert client.get("/api/dashboard").status_code == 401
    assert client.get("/api/dashboard", headers={"Authorization": "Bearer nope"}).status_code == 401
    assert client.post("/api/auth/login", data={"username": "admin@test.io", "password": "wrong"}).status_code == 401


def test_roles(client):
    h = login(client)
    assert client.post("/api/users", headers=h, json={"email": "v@test.io", "password": "viewer-password", "role": "viewer"}).status_code == 200
    hv = login(client, "v@test.io", "viewer-password")
    assert client.get("/api/dashboard", headers=hv).status_code == 200
    assert client.post("/api/websites", headers=hv, json={"domains": ["example.com"]}).status_code == 403
    assert client.put("/api/settings/qualification", headers=hv, json={"value": {}}).status_code == 403
    assert client.get("/api/audit", headers=hv).status_code == 403


def test_add_search_profile_export(client):
    h = login(client)
    r = client.post("/api/websites", headers=h, json={"domains": ["https://www.Example.com/a", "news.example.com", "example.com", "bad domain"]})
    body = r.json()
    assert body["added"] == 2 and len(body["invalid"]) == 1
    r = client.post("/api/websites/search", headers=h, json={"q": "example", "sort": ["host"]})
    assert r.json()["total"] == 2
    wid = r.json()["rows"][0]["id"]
    prof = client.get(f"/api/websites/{wid}", headers=h).json()
    assert prof["website"]["host"] == "example.com"
    assert prof["stages"][0]["stage"] == "qualify"
    assert prof["sources"][0]["type"] == "manual"

    r = client.post("/api/export", headers=h, json={"filter": {"q": "example"}, "format": "csv"})
    rows = list(csv.reader(io.StringIO(r.text)))
    assert rows[0][1] == "host" and sorted(x[1] for x in rows[1:]) == ["example.com", "news.example.com"]
    assert "not visits" in r.headers["x-data-notes"]
    r = client.post("/api/export", headers=h, json={"filter": {"q": "example"}, "format": "xlsx"})
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert wb["publishers"].max_row == 3


def test_override_requires_reason_and_is_audited(client):
    h = login(client)
    wid = client.post("/api/websites", headers=h, json={"domains": ["example.org"]}).json()["ids"][0]
    assert client.post(f"/api/websites/{wid}/override", headers=h, json={"field": "queue", "value": "VERIFIED_HIGH_VALUE", "reason": ""}).status_code == 422
    r = client.post(f"/api/websites/{wid}/override", headers=h, json={"field": "queue", "value": "VERIFIED_HIGH_VALUE", "reason": "traffic verified manually"})
    assert r.status_code == 200
    prof = client.get(f"/api/websites/{wid}", headers=h).json()
    assert prof["score"]["queue"] == "VERIFIED_HIGH_VALUE" and prof["score"]["computed_queue"] != "VERIFIED_HIGH_VALUE"
    assert any(a["action"] == "override" for a in client.get("/api/audit", headers=h).json())


def test_import_preview_and_job(client, tmp_path):
    h = login(client)
    data = "Website;Category\nexample.com;PRIORITY\nwww.example.com;UNKNOWN\nnot a site;\n"
    r = client.post("/api/imports/upload", headers=h, files={"file": ("x.csv", data, "text/csv")})
    assert r.status_code == 200
    up = r.json()
    assert up["suggested_mapping"]["domain"] == "Website" and up["suggested_mapping"]["legacy_label"] == "Category"
    r = client.post("/api/imports/start", headers=h, json={"token": up["token"], "mapping": up["suggested_mapping"]})
    jid = r.json()["job_id"]
    from app.jobs.base import run_job

    assert run_job(jid) == "DONE"
    job = next(j for j in client.get("/api/jobs", headers=h).json() if j["id"] == jid)
    assert job["progress"]["new_websites"] == 1 and job["progress"]["invalid"] == 1
    assert run_job(jid) == "DONE"  # idempotent


def test_job_resumes_from_checkpoint(clean_db, tmp_path):
    """Simulate a crash: checkpoint mid-file, then resume. No row is lost or duplicated."""
    from sqlalchemy import func, select, update

    from app.core.db import session_scope
    from app.jobs.base import create_job, run_job
    from app.models import Job, Website

    p = tmp_path / "d.csv"
    p.write_text("domain\n" + "\n".join(f"site{i}.com" for i in range(5000)))
    jid = create_job("import_file", {"path": str(p), "mapping": {"domain": "domain"}})
    assert run_job(jid) == "DONE"
    # pretend the worker died after row 2500: reset to RUNNING with an earlier checkpoint and resume
    with session_scope() as s:
        s.execute(update(Job).where(Job.id == jid).values(status="RUNNING", checkpoint={"row": 2500}))
    assert run_job(jid) == "DONE"
    with session_scope() as s:
        assert s.scalar(select(func.count()).select_from(Website)) == 5000


def test_supply_search_requires_param(client):
    h = login(client)
    assert client.get("/api/supply/search", headers=h).status_code == 400
